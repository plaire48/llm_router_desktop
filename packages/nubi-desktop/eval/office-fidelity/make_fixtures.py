# Office fidelity fixtures (사용자 지시 2026-10-21 "이미지, 도형, 글자, 폰트, 오피스에서 할 수 있는 모든 걸 다 테스트"):
# one synthetic Word, Excel and PowerPoint file each, holding as many Office features as the libraries and hand-written
# OOXML allow, so a capture engine can be graded feature by feature (features.mjs). Synthetic content only.
#
# Usage: python3 eval/office-fidelity/make_fixtures.py <out-dir>
# Requires python-docx 1.2+, openpyxl 3.1+, python-pptx 1.0+, Pillow. The generated files are committed under
# tests/fixtures/office-fidelity/ so tests and CI do not need Python.
import io, os, re, sys, zipfile, datetime
from PIL import Image, ImageDraw

OUT = sys.argv[1] if len(sys.argv) > 1 else '.'
os.makedirs(OUT, exist_ok=True)

def picture(color, text, size=(240, 120)):
    image = Image.new('RGB', size, color)
    draw = ImageDraw.Draw(image)
    draw.rectangle([4, 4, size[0] - 5, size[1] - 5], outline=(255, 255, 255), width=3)
    draw.text((16, size[1] // 2 - 6), text, fill=(255, 255, 255))
    path = os.path.join(OUT, f'.{text}.png')
    image.save(path)
    return path

LOGO = picture((31, 56, 100), 'LOGO')
PHOTO = picture((192, 0, 0), 'PHOTO', (300, 200))

# ---------------------------------------------------------------------------------------------------------- Word
def word():
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_UNDERLINE, WD_COLOR_INDEX, WD_BREAK
    from docx.enum.section import WD_SECTION, WD_ORIENT
    from docx.enum.dml import MSO_THEME_COLOR
    from docx.oxml import parse_xml
    from docx.oxml.ns import qn, nsdecls
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.part import Part
    from docx.opc.packuri import PackURI

    doc = Document()
    W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    R = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'

    def east(run, font):
        run.font.name = font
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), font)

    doc.add_heading('기능 점검 문서', level=0)
    doc.add_heading('1. 글자 서식', level=1)
    p = doc.add_paragraph()
    for text, font, size in [('맑은 고딕 11 ', '맑은 고딕', 11), ('Arial 14 ', 'Arial', 14), ('Times New Roman 9 ', 'Times New Roman', 9), ('궁서 20', '궁서', 20)]:
        run = p.add_run(text); east(run, font); run.font.size = Pt(size)
    p = doc.add_paragraph()
    r = p.add_run('굵게 '); r.bold = True
    r = p.add_run('기울임 '); r.italic = True
    r = p.add_run('밑줄 두 줄 '); r.underline = WD_UNDERLINE.DOUBLE
    r = p.add_run('취소선 '); r.font.strike = True
    r = p.add_run('이중 취소선 '); r.font.double_strike = True
    r = p.add_run('위 첨자'); r.font.superscript = True
    p.add_run(' ')
    r = p.add_run('아래 첨자'); r.font.subscript = True
    p.add_run(' ')
    r = p.add_run('Small Caps '); r.font.small_caps = True
    r = p.add_run('all caps '); r.font.all_caps = True
    r = p.add_run('빨간 글자 '); r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    r = p.add_run('테마 색 '); r.font.color.theme_color = MSO_THEME_COLOR.ACCENT_1
    r = p.add_run('형광펜 '); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    r = p.add_run('자간 넓게 '); r._element.get_or_add_rPr().append(parse_xml(f'<w:spacing {W} w:val="40"/>'))
    r = p.add_run('글자 음영'); r._element.get_or_add_rPr().append(parse_xml(f'<w:shd {W} w:val="clear" w:color="auto" w:fill="D9E1F2"/>'))
    p = doc.add_paragraph('기호와 이모지: ★ ✓ ① ㉠ ※ 😀 € ¥')

    doc.add_heading('2. 문단 서식', level=1)
    for text, align in [('가운데 정렬 문단', WD_ALIGN_PARAGRAPH.CENTER), ('오른쪽 정렬 문단', WD_ALIGN_PARAGRAPH.RIGHT), ('양쪽 정렬 문단은 줄 끝을 맞춘다. ' * 3, WD_ALIGN_PARAGRAPH.JUSTIFY), ('배분 정렬', WD_ALIGN_PARAGRAPH.DISTRIBUTE)]:
        doc.add_paragraph(text).alignment = align
    p = doc.add_paragraph('들여쓰기 2cm, 첫 줄 1cm, 앞 12pt 뒤 6pt, 줄 간격 1.5')
    f = p.paragraph_format; f.left_indent = Cm(2); f.first_line_indent = Cm(1); f.space_before = Pt(12); f.space_after = Pt(6); f.line_spacing = 1.5; f.keep_with_next = True
    p = doc.add_paragraph('내어쓰기 문단은 둘째 줄부터 왼쪽에 붙는다. ' * 2); p.paragraph_format.left_indent = Cm(1); p.paragraph_format.first_line_indent = Cm(-1)
    p = doc.add_paragraph('테두리와 음영이 있는 문단')
    p._element.get_or_add_pPr().append(parse_xml(f'<w:pBdr {W}><w:top w:val="single" w:sz="8" w:space="1" w:color="1F3864"/><w:bottom w:val="double" w:sz="6" w:space="1" w:color="1F3864"/></w:pBdr>'))
    p._element.get_or_add_pPr().append(parse_xml(f'<w:shd {W} w:val="clear" w:color="auto" w:fill="FFF2CC"/>'))
    p = doc.add_paragraph('탭\t다음 글자와 줄 바꿈'); p.add_run().add_break(WD_BREAK.LINE); p.add_run('둘째 줄')

    doc.add_heading('3. 목록', level=1)
    doc.add_paragraph('글머리 기호 1', style='List Bullet')
    doc.add_paragraph('글머리 기호 2단계', style='List Bullet 2')
    doc.add_paragraph('번호 목록 1', style='List Number')
    doc.add_paragraph('번호 목록 2', style='List Number')
    doc.add_paragraph('번호 목록 2단계', style='List Number 2')

    doc.add_heading('4. 표', level=1)
    table = doc.add_table(rows=4, cols=4); table.style = 'Table Grid'
    for c, label in enumerate(['구분', '1분기', '2분기', '합계']):
        cell = table.cell(0, c); cell.text = label; cell.paragraphs[0].runs[0].bold = True
        cell._element.get_or_add_tcPr().append(parse_xml(f'<w:shd {W} w:val="clear" w:color="auto" w:fill="D9E1F2"/>'))
    table.cell(1, 0).text = '매출'; table.cell(1, 1).text = '100'; table.cell(1, 2).text = '120'; table.cell(1, 3).text = '220'
    table.cell(2, 0).merge(table.cell(3, 0)).text = '세로 병합'
    table.cell(2, 1).merge(table.cell(2, 2)).text = '가로 병합'
    table.cell(2, 3).text = '비고'
    inner = table.cell(3, 3).add_table(rows=2, cols=2)
    inner.cell(0, 0).text = '중첩 표'; inner.cell(1, 1).text = '안쪽 칸'
    tr = table.rows[0]._tr; tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {W}/>'))

    doc.add_heading('5. 그림과 도형', level=1)
    doc.add_paragraph('인라인 그림:')
    doc.add_picture(LOGO, width=Cm(4))
    rid, _ = doc.part.get_or_add_image(PHOTO)
    WP = 'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture"'
    anchor = f'''<w:r {W} {R}><w:drawing><wp:anchor {WP} distT="0" distB="0" distL="114300" distR="114300" simplePos="0" relativeHeight="251659264" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">
<wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="column"><wp:posOffset>3600000</wp:posOffset></wp:positionH><wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV>
<wp:extent cx="1800000" cy="1200000"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapSquare wrapText="bothSides"/><wp:docPr id="50" name="떠 있는 사진" descr="빨간 사진"/><wp:cNvGraphicFramePr/>
<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic><pic:nvPicPr><pic:cNvPr id="50" name="떠 있는 사진"/><pic:cNvPicPr/></pic:nvPicPr>
<pic:blipFill><a:blip r:embed="{rid}"/><a:srcRect l="5000" t="5000" r="5000" b="5000"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill><pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1800000" cy="1200000"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic></a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>'''
    p = doc.add_paragraph('떠 있는 그림(오른쪽 배치)이 이 문단에 붙어 있다.'); p._element.append(parse_xml(anchor))
    WPS = 'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
    def shape(id_, name, prst, fill, text, x):
        body = f'<wps:txbx><w:txbxContent><w:p><w:r><w:rPr><w:b/><w:color w:val="FFFFFF"/></w:rPr><w:t>{text}</w:t></w:r></w:p></w:txbxContent></wps:txbx>' if text else ''
        return f'''<w:r {W} {R}><w:drawing><wp:anchor {WP} distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="251660{id_}" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">
<wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="column"><wp:posOffset>{x}</wp:posOffset></wp:positionH><wp:positionV relativeFrom="paragraph"><wp:posOffset>100000</wp:posOffset></wp:positionV>
<wp:extent cx="1600000" cy="700000"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapTopAndBottom/><wp:docPr id="{id_}" name="{name}"/><wp:cNvGraphicFramePr/>
<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp {WPS}><wps:cNvSpPr/><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1600000" cy="700000"/></a:xfrm><a:prstGeom prst="{prst}"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="{fill}"/></a:solidFill><a:ln w="19050"><a:solidFill><a:srgbClr val="000000"/></a:solidFill><a:prstDash val="dash"/></a:ln></wps:spPr>{body}<wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>'''
    p = doc.add_paragraph('도형 문단'); p._element.append(parse_xml(shape(61, '글상자', 'rect', '2E75B6', '글상자 안 글', 0))); p._element.append(parse_xml(shape(62, '둥근 사각형', 'roundRect', '70AD47', '', 1900000))); p._element.append(parse_xml(shape(63, '오른쪽 화살표', 'rightArrow', 'ED7D31', '', 3800000)))

    doc.add_heading('6. 차트와 수식', level=1)
    from pptx.chart.xmlwriter import ChartXmlWriter
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    data = CategoryChartData(); data.categories = ['1분기', '2분기', '3분기']; data.add_series('매출', (100, 120, 90))
    chart_xml = ChartXmlWriter(XL_CHART_TYPE.COLUMN_CLUSTERED, data).xml
    chart_xml = re.sub(r'<c:autoTitleDeleted val="\d"/>', '', chart_xml)
    chart_xml = chart_xml.replace('<c:plotArea>', '<c:title><c:tx><c:rich><a:bodyPr/><a:p><a:r><a:t>분기 매출</a:t></a:r></a:p></c:rich></c:tx><c:overlay val="0"/></c:title><c:autoTitleDeleted val="0"/><c:plotArea>', 1)
    chart_part = Part(PackURI('/word/charts/chart1.xml'), 'application/vnd.openxmlformats-officedocument.drawingml.chart+xml', chart_xml.encode('utf-8'), doc.part.package)
    crid = doc.part.relate_to(chart_part, RT.CHART)
    C = 'xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"'
    p = doc.add_paragraph(); p._element.append(parse_xml(f'''<w:r {W} {R}><w:drawing><wp:inline {WP}><wp:extent cx="4500000" cy="2500000"/><wp:docPr id="70" name="차트 1"/><a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:chart {C} r:id="{crid}"/></a:graphicData></a:graphic></wp:inline></w:drawing></w:r>'''))
    M = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'
    p = doc.add_paragraph('수식: '); p._element.append(parse_xml(f'<m:oMath {M} {W}><m:r><m:t>E=m</m:t></m:r><m:sSup><m:e><m:r><m:t>c</m:t></m:r></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup></m:oMath>'))

    doc.add_heading('7. 링크, 책갈피, 필드, 주석', level=1)
    hrid = doc.part.relate_to('https://example.com/form-guide', RT.HYPERLINK, is_external=True)
    p = doc.add_paragraph('바깥 링크: '); p._element.append(parse_xml(f'<w:hyperlink {W} {R} r:id="{hrid}"><w:r><w:rPr><w:rStyle w:val="Hyperlink"/><w:color w:val="0563C1"/><w:u w:val="single"/></w:rPr><w:t>작성 안내</w:t></w:r></w:hyperlink>'))
    p = doc.add_paragraph(); p._element.append(parse_xml(f'<w:bookmarkStart {W} w:id="10" w:name="결재란"/>')); p.add_run('책갈피가 걸린 문단'); p._element.append(parse_xml(f'<w:bookmarkEnd {W} w:id="10"/>'))
    p = doc.add_paragraph('오늘 날짜 필드: ')
    for xml in [f'<w:r {W}><w:fldChar w:fldCharType="begin"/></w:r>', f'<w:r {W}><w:instrText xml:space="preserve"> DATE \\@ "yyyy-MM-dd" </w:instrText></w:r>', f'<w:r {W}><w:fldChar w:fldCharType="separate"/></w:r>', f'<w:r {W}><w:t>2026-10-21</w:t></w:r>', f'<w:r {W}><w:fldChar w:fldCharType="end"/></w:r>']:
        p._element.append(parse_xml(xml))
    p = doc.add_paragraph('각주가 붙는 문장'); p._element.append(parse_xml(f'<w:r {W}><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteReference w:id="1"/></w:r>'))
    p.add_run(' 그리고 미주'); p._element.append(parse_xml(f'<w:r {W}><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:endnoteReference w:id="1"/></w:r>'))
    def notes(kind, text):
        tag = 'footnote' if kind == 'footnotes' else 'endnote'
        return f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:{kind} {W}><w:{tag} w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:{tag}><w:{tag} w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:{tag}><w:{tag} w:id="1"><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:{tag}></w:{kind}>'''.encode('utf-8')
    doc.part.relate_to(Part(PackURI('/word/footnotes.xml'), 'application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml', notes('footnotes', '각주 내용: 출처는 합성 자료.'), doc.part.package), RT.FOOTNOTES)
    doc.part.relate_to(Part(PackURI('/word/endnotes.xml'), 'application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml', notes('endnotes', '미주 내용: 참고 문헌.'), doc.part.package), RT.ENDNOTES)
    p = doc.add_paragraph(); run = p.add_run('검토 의견이 달린 글')
    doc.add_comment(run, text='이 문장은 확인이 필요합니다.', author='검토자', initials='검')
    p = doc.add_paragraph('변경 추적: ')
    p._element.append(parse_xml(f'<w:ins {W} w:id="101" w:author="편집자" w:date="2026-10-21T09:00:00Z"><w:r><w:t>삽입된 글</w:t></w:r></w:ins>'))
    p._element.append(parse_xml(f'<w:del {W} w:id="102" w:author="편집자" w:date="2026-10-21T09:01:00Z"><w:r><w:delText>지운 글</w:delText></w:r></w:del>'))

    doc.add_heading('8. 입력 컨트롤', level=1)
    W14 = 'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml"'
    def sdt(pr, text):
        return parse_xml(f'<w:sdt {W} {W14}><w:sdtPr>{pr}</w:sdtPr><w:sdtContent><w:r><w:t>{text}</w:t></w:r></w:sdtContent></w:sdt>')
    p = doc.add_paragraph('부서: '); p._element.append(sdt('<w:alias w:val="부서명"/><w:tag w:val="dept"/><w:showingPlcHdr/><w:text/>', '부서를 입력하세요'))
    p = doc.add_paragraph('작성일: '); p._element.append(sdt('<w:alias w:val="작성일"/><w:tag w:val="date"/><w:date w:fullDate="2026-10-21T00:00:00Z"><w:dateFormat w:val="yyyy-MM-dd"/><w:lid w:val="ko-KR"/></w:date>', '2026-10-21'))
    p = doc.add_paragraph('구분: '); p._element.append(sdt('<w:alias w:val="구분"/><w:tag w:val="kind"/><w:dropDownList><w:listItem w:displayText="신규" w:value="new"/><w:listItem w:displayText="변경" w:value="change"/></w:dropDownList>', '신규'))
    p = doc.add_paragraph('확인: '); p._element.append(sdt('<w:alias w:val="확인"/><w:tag w:val="ok"/><w14:checkbox><w14:checked w14:val="1"/><w14:checkedState w14:val="2612"/><w14:uncheckedState w14:val="2610"/></w14:checkbox>', '☒'))
    p = doc.add_paragraph('쪽 나눔 뒤 새 구역'); p.add_run().add_break(WD_BREAK.PAGE)

    first = doc.sections[0]
    first.different_first_page_header_footer = True
    first.header.paragraphs[0].text = '기본 머리글'
    first.first_page_header.paragraphs[0].text = '첫 쪽 머리글'
    fp = first.footer.paragraphs[0]; fp.text = '쪽 '; fp._element.append(parse_xml(f'<w:fldSimple {W} w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>'))
    second = doc.add_section(WD_SECTION.NEW_PAGE)
    second.orientation = WD_ORIENT.LANDSCAPE; second.page_width, second.page_height = first.page_height, first.page_width
    cols = second._sectPr.find(qn('w:cols'))
    if cols is None:
        cols = parse_xml(f'<w:cols {W}/>'); second._sectPr.append(cols)
    cols.set(qn('w:num'), '2'); cols.set(qn('w:space'), '720')
    doc.add_heading('9. 가로 2단 구역', level=1)
    doc.add_paragraph('가로 방향, 2단으로 나뉜 구역의 본문. ' * 12)
    doc.save(os.path.join(OUT, 'synthetic-kitchen-sink.docx'))

# ---------------------------------------------------------------------------------------------------------- Excel
def excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, GradientFill, Border, Side, Alignment
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, IconSetRule
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.comments import Comment
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.formula import ArrayFormula
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont

    wb = Workbook(); ws = wb.active; ws.title = '요약'
    ws.sheet_properties.tabColor = '1F3864'
    ws['A1'] = '기능 점검 통합 문서'; ws['A1'].font = Font(name='맑은 고딕', size=18, bold=True, color='1F3864'); ws.merge_cells('A1:F1')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center'); ws.row_dimensions[1].height = 32
    styles = [('굵게', Font(bold=True)), ('기울임', Font(italic=True)), ('밑줄', Font(underline='double')), ('취소선', Font(strike=True)), ('빨강', Font(color='C00000')), ('Arial 14', Font(name='Arial', size=14)), ('위 첨자', Font(vertAlign='superscript'))]
    for i, (text, font) in enumerate(styles):
        cell = ws.cell(row=3, column=i + 1, value=text); cell.font = font
    ws['A4'] = '단색 채움'; ws['A4'].fill = PatternFill('solid', fgColor='FFF2CC')
    ws['B4'] = '무늬 채움'; ws['B4'].fill = PatternFill('darkGrid', fgColor='9BC2E6', bgColor='FFFFFF')
    ws['C4'] = '그라데이션'; ws['C4'].fill = GradientFill(stop=('FF0000', '0000FF'))
    thin, thick, dashed = Side(style='thin', color='000000'), Side(style='thick', color='C00000'), Side(style='dashed', color='2E75B6')
    ws['D4'] = '테두리'; ws['D4'].border = Border(top=thick, bottom=dashed, left=thin, right=thin)
    ws['E4'] = '45도 회전'; ws['E4'].alignment = Alignment(textRotation=45)
    ws['F4'] = '줄 바꿈 되는 긴 글자입니다'; ws['F4'].alignment = Alignment(wrap_text=True, vertical='top', indent=1)
    ws['A5'] = CellRichText(['일부만 ', TextBlock(InlineFont(b=True, color='C00000'), '빨간 굵은'), ' 글자'])
    ws['A6'] = '메모가 달린 칸'; ws['A6'].comment = Comment('합성 메모: 확인 필요', '검토자')
    ws['A7'] = '바깥 링크'; ws['A7'].hyperlink = 'https://example.com/sheet-guide'
    ws['B7'] = '안쪽 링크'; ws['B7'].hyperlink = '#데이터!A1'
    ws['A8'] = 0.256; ws['A8'].number_format = '0.0%'
    ws['B8'] = 1250000; ws['B8'].number_format = '"₩"#,##0'
    ws['C8'] = datetime.date(2026, 10, 21); ws['C8'].number_format = 'yyyy-mm-dd'
    ws['D8'] = True
    ws['E8'] = 3.14159; ws['E8'].number_format = '0.00_);(0.00)'
    ws.column_dimensions['A'].width = 18; ws.column_dimensions['F'].width = 22
    ws.column_dimensions['G'].hidden = True; ws['G2'] = '숨긴 열'
    ws.row_dimensions[9].hidden = True; ws['A9'] = '숨긴 행'
    ws.freeze_panes = 'A3'
    ws.add_image(XLImage(LOGO), 'H2')

    data = wb.create_sheet('데이터')
    data.append(['품목', '수량', '단가', '금액'])
    for row in [('사무용 의자', 10, 120000), ('모니터', 5, 280000), ('키보드', 20, 35000)]:
        data.append([*row, None])
    for r in range(2, 5):
        data[f'D{r}'] = f'=B{r}*C{r}'; data[f'C{r}'].number_format = '#,##0'; data[f'D{r}'].number_format = '#,##0'
    data['A5'] = '합계'; data['D5'] = '=SUM(D2:D4)'; data['D5'].font = Font(bold=True)
    data['F2'] = ArrayFormula('F2:F4', '=B2:B4*2')
    data['A7'] = '요약 참조'; data['B7'] = "='요약'!B8*2"
    data['A8'] = '찾기'; data['B8'] = '=VLOOKUP("모니터",A2:D4,4,FALSE)'
    table = Table(displayName='SalesTable', ref='A1:D4'); table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True); data.add_table(table)
    data.conditional_formatting.add('B2:B4', CellIsRule(operator='greaterThan', formula=['8'], fill=PatternFill('solid', fgColor='FFC7CE')))
    data.conditional_formatting.add('C2:C4', ColorScaleRule(start_type='min', start_color='F8696B', end_type='max', end_color='63BE7B'))
    data.conditional_formatting.add('D2:D4', DataBarRule(start_type='min', end_type='max', color='638EC6'))
    data.conditional_formatting.add('B2:B4', IconSetRule('3Arrows', 'percent', [0, 33, 67]))
    dv = DataValidation(type='list', formula1='"현금,카드,계좌"', allow_blank=True); data.add_data_validation(dv); dv.add('E2'); data['E1'] = '결제'
    dv2 = DataValidation(type='whole', operator='between', formula1='1', formula2='100'); data.add_data_validation(dv2); dv2.add('B2:B4')
    data.row_dimensions.group(2, 4, outline_level=1)
    data.protection.sheet = True
    data.print_area = 'A1:D5'; data.page_setup.orientation = 'landscape'; data.print_title_rows = '1:1'
    wb.defined_names['합계금액'] = DefinedName('합계금액', attr_text="'데이터'!$D$5")

    charts = wb.create_sheet('차트')
    bar = BarChart(); bar.title = '품목별 금액'; bar.add_data(Reference(data, min_col=4, min_row=1, max_row=4), titles_from_data=True); bar.set_categories(Reference(data, min_col=1, min_row=2, max_row=4)); charts.add_chart(bar, 'A1')
    line = LineChart(); line.title = '수량 추이'; line.add_data(Reference(data, min_col=2, min_row=1, max_row=4), titles_from_data=True); charts.add_chart(line, 'J1')
    pie = PieChart(); pie.title = '수량 비율'; pie.add_data(Reference(data, min_col=2, min_row=1, max_row=4), titles_from_data=True); pie.set_categories(Reference(data, min_col=1, min_row=2, max_row=4)); charts.add_chart(pie, 'A16')
    hidden = wb.create_sheet('숨김'); hidden['A1'] = '숨긴 시트'; hidden.sheet_state = 'hidden'
    very = wb.create_sheet('아주숨김'); very['A1'] = '아주 숨긴 시트'; very.sheet_state = 'veryHidden'
    path = os.path.join(OUT, 'synthetic-kitchen-sink.xlsx'); wb.save(path)
    # A text box and an arrow on the summary sheet's drawing (openpyxl writes pictures and charts only).
    XDR = 'xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    def sp(id_, name, prst, fill, text, col):
        body = f'<xdr:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="ko-KR" b="1"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a:rPr><a:t>{text}</a:t></a:r></a:p></xdr:txBody>' if text else ''
        return f'<xdr:twoCellAnchor {XDR}><xdr:from><xdr:col>{col}</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>11</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:to><xdr:col>{col + 2}</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>14</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to><xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="{id_}" name="{name}"/><xdr:cNvSpPr txBox="{1 if text else 0}"/></xdr:nvSpPr><xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1200000" cy="500000"/></a:xfrm><a:prstGeom prst="{prst}"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="{fill}"/></a:solidFill></xdr:spPr>{body}</xdr:sp><xdr:clientData/></xdr:twoCellAnchor>'
    rewrite(path, lambda name, text: text.replace('</xdr:wsDr>', sp(90, '글상자 1', 'rect', '2E75B6', '시트 위 글상자', 0) + sp(91, '화살표 1', 'rightArrow', 'ED7D31', '', 3) + '</xdr:wsDr>') if re.match(r'xl/drawings/drawing\d+\.xml$', name) and 'pic:pic' in text else text)

def rewrite(path, change):
    with zipfile.ZipFile(path) as source:
        entries = [(info, source.read(info.filename)) for info in source.infolist()]
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as target:
        for info, data in entries:
            if info.filename.endswith('.xml') or info.filename.endswith('.rels'):
                data = change(info.filename, data.decode('utf-8')).encode('utf-8')
            target.writestr(info, data)

# ---------------------------------------------------------------------------------------------------------- PowerPoint
def powerpoint():
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls, qn

    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = '기능 점검 발표'
    r = s1.shapes.title.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(44); r.font.bold = True; r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64); r.font.name = '맑은 고딕'
    s1.placeholders[1].text = '합성 자료로 만든 점검용 슬라이드'
    s1.background.fill.solid(); s1.background.fill.fore_color.rgb = RGBColor(0xF2, 0xF2, 0xF2)

    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = '글머리와 수준'
    body = s2.placeholders[1].text_frame; body.text = '첫째 수준'
    for text, level in [('둘째 수준', 1), ('셋째 수준', 2), ('다시 첫째', 0)]:
        p = body.add_paragraph(); p.text = text; p.level = level
    s2.notes_slide.notes_text_frame.text = '발표자 노트: 목록 수준을 설명한다.'

    s3 = prs.slides.add_slide(prs.slide_layouts[6])
    box = s3.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(8), Inches(1.6)); tf = box.text_frame
    p = tf.paragraphs[0]
    for text, setup in [('굵게 ', lambda f: setattr(f, 'bold', True)), ('기울임 ', lambda f: setattr(f, 'italic', True)), ('밑줄 ', lambda f: setattr(f, 'underline', True)), ('Arial 24 ', lambda f: (setattr(f, 'name', 'Arial'), setattr(f, 'size', Pt(24)))), ('빨강', lambda f: setattr(f.color, 'rgb', RGBColor(0xC0, 0, 0)))]:
        run = p.add_run(); run.text = text; setup(run.font)
    run = p.add_run(); run.text = ' 취소선'; run.font._rPr.set('strike', 'sngStrike')
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER; link = p2.add_run(); link.text = '바깥 링크'; link.hyperlink.address = 'https://example.com/deck-guide'
    shapes = [(MSO_SHAPE.RECTANGLE, '사각형', '2E75B6', '사각형 안 글'), (MSO_SHAPE.ROUNDED_RECTANGLE, '둥근 사각형', None, '그라데이션'), (MSO_SHAPE.OVAL, '타원', '70AD47', ''), (MSO_SHAPE.RIGHT_ARROW, '오른쪽 화살표', 'ED7D31', ''), (MSO_SHAPE.STAR_5_POINT, '별', 'FFC000', '')]
    made = []
    for i, (kind, name, fill, text) in enumerate(shapes):
        shape = s3.shapes.add_shape(kind, Inches(0.5 + i * 2.5), Inches(2.3), Inches(2.2), Inches(1.2)); shape.name = name
        if fill:
            shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        else:
            shape.fill.gradient(); shape.fill.gradient_angle = 90
        shape.line.color.rgb = RGBColor(0, 0, 0); shape.line.width = Pt(2); shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        if text: shape.text_frame.text = text
        made.append(shape)
    group = s3.shapes.add_group_shape(); group.name = '묶음'
    g1 = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(4.0), Inches(1.5), Inches(0.8)); g1.name = '묶음 사각형'; g1.text_frame.text = '묶음 안 1'
    g2 = group.shapes.add_shape(MSO_SHAPE.OVAL, Inches(2.2), Inches(4.0), Inches(1.5), Inches(0.8)); g2.name = '묶음 타원'; g2.text_frame.text = '묶음 안 2'
    connector = s3.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(4.0), Inches(4.4), Inches(6.0), Inches(4.4)); connector.name = '연결선'
    connector.line._get_or_add_ln().append(parse_xml(f'<a:tailEnd {nsdecls("a")} type="triangle"/>'))
    art = s3.shapes.add_textbox(Inches(7), Inches(4.0), Inches(5), Inches(1.2)); art.name = '워드아트'; art.text_frame.text = '휘어진 글자'
    art.text_frame._txBody.find(qn('a:bodyPr')).append(parse_xml(f'<a:prstTxWarp {nsdecls("a")} prst="textArchUp"><a:avLst/></a:prstTxWarp>'))
    pic = s3.shapes.add_picture(PHOTO, Inches(0.5), Inches(5.2), height=Inches(1.8)); pic.name = '잘린 사진'; pic.crop_left = 0.1; pic.crop_right = 0.1
    timing = '<p:timing ' + nsdecls('p') + '><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst><p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst><p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:par><p:cTn id="5" presetID="1" presetClass="entr" presetSubtype="0" fill="hold" nodeType="clickEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:set><p:cBhvr><p:cTn id="6" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="' + str(made[0].shape_id) + '"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'
    s3._element.append(parse_xml(f'<p:transition {nsdecls("p")} spd="med"><p:fade/></p:transition>'))
    s3._element.append(parse_xml(timing))

    s4 = prs.slides.add_slide(prs.slide_layouts[5]); s4.shapes.title.text = '표와 차트'
    table = s4.shapes.add_table(4, 3, Inches(0.5), Inches(1.5), Inches(5.5), Inches(2.4)).table
    for c, label in enumerate(['구분', '상반기', '하반기']): table.cell(0, c).text = label
    table.cell(1, 0).text = '매출'; table.cell(1, 1).text = '100'; table.cell(1, 2).text = '120'
    table.cell(2, 0).merge(table.cell(3, 0)); table.cell(2, 0).text = '세로 병합'
    table.cell(2, 1).merge(table.cell(2, 2)); table.cell(2, 1).text = '가로 병합'
    for kind, title, x in [(XL_CHART_TYPE.COLUMN_CLUSTERED, '분기 매출', 6.3), (XL_CHART_TYPE.PIE, '비율', 9.6)]:
        data = CategoryChartData(); data.categories = ['1분기', '2분기', '3분기']; data.add_series('매출', (100, 120, 90))
        chart = s4.shapes.add_chart(kind, Inches(x), Inches(1.5), Inches(3.2), Inches(2.6), data).chart; chart.has_title = True; chart.chart_title.text_frame.text = title
    mp4 = os.path.join(OUT, '.clip.mp4')
    with open(mp4, 'wb') as handle: handle.write(b'\x00\x00\x00\x18ftypmp42\x00\x00\x00\x00mp42isom' + b'\x00' * 64)
    s4.shapes.add_movie(mp4, Inches(0.5), Inches(4.3), Inches(3), Inches(1.7), poster_frame_image=LOGO, mime_type='video/mp4')

    s5 = prs.slides.add_slide(prs.slide_layouts[6]); s5._element.set('show', '0')
    s5.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = '숨긴 슬라이드'
    ids = [str(slide.slide_id) for slide in prs.slides]
    ext = f'<p:extLst {nsdecls("p")}><p:ext uri="{{521415D9-36F7-43E2-AB2F-B90AF26B5E84}}"><p14:sectionLst xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main"><p14:section name="소개" id="{{11111111-1111-1111-1111-111111111111}}"><p14:sldIdLst>' + ''.join(f'<p14:sldId id="{i}"/>' for i in ids[:2]) + '</p14:sldIdLst></p14:section><p14:section name="본문" id="{{22222222-2222-2222-2222-222222222222}}"><p14:sldIdLst>' + ''.join(f'<p14:sldId id="{i}"/>' for i in ids[2:]) + '</p14:sldIdLst></p14:section></p14:sectionLst></p:ext></p:extLst>'
    prs.part._element.append(parse_xml(ext))
    prs.save(os.path.join(OUT, 'synthetic-kitchen-sink.pptx'))

word(); excel(); powerpoint()
for name in os.listdir(OUT):
    if name.startswith('.') and (name.endswith('.png') or name.endswith('.mp4')): os.remove(os.path.join(OUT, name))
print('fixtures written to', OUT)
