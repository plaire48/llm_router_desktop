"""Independent saved-PPTX/embedded-data check for the reviewed final candidates."""
import hashlib,json
from pathlib import Path
from zipfile import ZipFile
from io import BytesIO
from lxml import etree as E
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];ns={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
paths=['quality-delivery-6/codex-delivery-powerpoint','quality-observed-pages-5/claude-inventory','quality-consensus-3/claude-delivery-powerpoint','quality-minor-gate-4/codex-renewal-powerpoint','quality-observed-pages-5/claude-renewal-powerpoint'];rows=[]
for relative in paths:
 d=ROOT/'artifacts'/relative;r=json.loads((d/'result.json').read_text());raw=(d/'final.pptx').read_bytes();sha=hashlib.sha256(raw).hexdigest();receipts=[json.loads(p.read_text()) for p in d.glob('round-*/receipt.json')];receipt=next(x for x in receipts if x['artifactHash']==sha);matching=next(p.parent for p in d.glob('round-*/receipt.json') if json.loads(p.read_text())['artifactHash']==sha);reviews={p.name.split('-')[1]:[] for p in matching.glob('visual-*.json') if '-attempt-' not in p.name}
 for provider in reviews:
  reviews[provider]=[json.loads(p.read_text()) for p in matching.glob('visual-'+provider+'-*.json') if '-attempt-' not in p.name]
 series_count=0;count=receipt['pages']['count'];coverage={provider:sorted({page for rv in group for page in rv['reviewedPages']}) for provider,group in reviews.items()};checks=[]
 with ZipFile(BytesIO(raw)) as z:
  slides=[p for p in z.namelist() if p.startswith('ppt/slides/slide') and p.endswith('.xml') and '/_rels/' not in p];tables=sum(z.read(p).count(b'<a:tbl>') for p in slides);charts=[p for p in z.namelist() if p.startswith('ppt/charts/chart') and p.endswith('.xml') and '/_rels/' not in p]
  for chart in charts:
   n=Path(chart).stem.replace('chart','');book=load_workbook(BytesIO(z.read('ppt/embeddings/NubiChartData'+n+'.xlsx')),data_only=True);tree=E.fromstring(z.read(chart));series=tree.findall('.//c:ser',ns);series_count+=len(series)
   for ci,s in enumerate(series):
    values=[float(p.find('c:v',ns).text) for p in s.findall('c:val/c:numRef/c:numCache/c:pt',ns)];actual=[book['Sheet1'].cell(i+2,ci+2).value for i in range(len(values))];checks.append(values==actual)
   if tree.find('.//c:barChart',ns) is not None:
    negatives=any(float(v.text)<0 for v in tree.findall('.//c:val//c:pt/c:v',ns))
    if negatives:checks.extend([all(s.find('c:invertIfNegative',ns) is not None and s.find('c:invertIfNegative',ns).get('val')=='0' for s in series),tree.find('.//c:catAx/c:tickLblPos',ns).get('val')=='low'])
 passed=r['status']=='verified_draft' and len(slides)==count and tables>0 and charts and all(checks) and all(c['status']=='passed' for c in receipt['checks']) and set(coverage)=={'claude','codex'} and all(pages==list(range(1,count+1)) for pages in coverage.values())
 rows.append(dict(path=relative,artifactHash=sha,passed=bool(passed),pages=count,nativeTables=tables,editableCharts=len(charts),chartSeries=series_count,matchingChartChecks=sum(checks),reviewCoverage=coverage))
(ROOT/'artifacts/quality-final-presentations.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2)+'\n');assert all(r['passed'] for r in rows),rows;print(json.dumps({'passed':sum(r['passed'] for r in rows),'total':len(rows),'pages':sum(r['pages'] for r in rows)}))
