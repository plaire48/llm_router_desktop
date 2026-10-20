"""Independent source arithmetic + actual LibreOffice calculation and rendered file QA.
Only newly generated synthetic drafts are opened. No model-auth or native Office access.
"""
import sys,json,os,subprocess,tempfile,re,copy
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree as E
from openpyxl import load_workbook
from pypdf import PdfReader
from pdf2image import convert_from_path
ROOT=Path(sys.argv[1]).resolve(); FILTER=sys.argv[2] if len(sys.argv)>2 and not sys.argv[2].startswith('--') else None
RUNTIME=Path('/Users/ganadi/.cache/codex-runtimes/codex-primary-runtime/dependencies');BIN=RUNTIME/'bin/override';QA=ROOT/'qa';QA.mkdir(exist_ok=True)
FONT=QA/'fonts';FONT.mkdir(exist_ok=True);CONFIG=FONT/'fonts.conf';CONFIG.write_text(f'<fontconfig><dir>/System/Library/Fonts</dir><dir>/System/Library/Fonts/Supplemental</dir><cachedir>{FONT}/cache</cachedir></fontconfig>');os.environ.update(FONTCONFIG_FILE=str(CONFIG),SAL_FONTPATH='/System/Library/Fonts',TMPDIR='/private/tmp',PATH=str(BIN)+os.pathsep+os.environ.get('PATH',''))
NS={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
cases=json.loads(subprocess.run([str(RUNTIME/'node/bin/node'),'--input-type=module','-e',"import{cases}from'./eval/dual-business-cases.mjs';console.log(JSON.stringify(cases))"],capture_output=True,text=True,check=True).stdout)
def expected(c):
 d=c['datasets'];rnd=lambda x:round(x+1e-10,2)
 if c['id']=='maintenance':
  accepted=[x for x in d['jobs'] if x['state']=='승인'];material=sum(x['qty']*x['price'] for x in accepted);labor=sum(x['hours']*x['rate'] for x in accepted);direct=material+labor;s=d['settings'][0];vat=rnd(material*s['vat']);reserve=rnd(direct*s['reserve']);total=rnd(direct+vat+reserve)
  return dict(material=material,labor=labor,approved=direct,vat=vat,reserve=reserve,total=total,gap=rnd(s['budget']-total),pending=sum(x['qty']*x['price']+x['hours']*x['rate'] for x in d['jobs'] if x['state']=='보류'))
 if c['id']=='inventory':
  available=replenish=purchase=value=negative=0
  for x in d['stock']:
   ending=x['opening']+sum(m['quantity'] for m in d['movements'] if m['warehouse']==x['warehouse'] and m['sku']==x['sku'] and m['condition']=='정상' and m['state']=='유효');need=max(0,x['target']-ending);available+=ending;replenish+=need;purchase+=need*x['unitCost'];value+=ending*x['unitCost'];negative+=ending<0
  return dict(available=available,quarantined=sum(x['quantity'] for x in d['movements'] if x['condition']=='격리' and x['type']=='반품' and x['state']=='유효'),replenish=replenish,purchase=purchase,value=value,negative=negative,cancelled=sum(x['state']=='취소' for x in d['movements']))
 closed=[x for x in d['tickets'] if x['state']=='완료'];ongoing=[x for x in d['tickets'] if x['state']=='진행'];met=sum(x['elapsed']<=x['target'] for x in closed);work=sum(x['work'] for x in closed)+sum(x['remaining'] for x in ongoing);capacity=sum(x['people']*x['minutes'] for x in d['capacity'])
 return dict(closed=len(closed),met=met,breached=len(closed)-met,sla=met/len(closed) if closed else 0,credit=rnd(sum(x['fee']*x['creditRate'] for x in closed if x['elapsed']>x['target'])),potential=rnd(sum(x['fee']*x['creditRate'] for x in ongoing if x['elapsed']>x['target'])),work=work,capacity=capacity,gap=capacity-work)
def variants(c):
 change=copy.deepcopy(c);zero=copy.deepcopy(c)
 if c['id']=='maintenance':
  change['datasets']['jobs'][1]['price']+=17
  for x in zero['datasets']['jobs']:x['qty']=x['hours']=0
 elif c['id']=='inventory':
  change['datasets']['movements'][1]['quantity']+=15
  for x in zero['datasets']['movements']:x['quantity']=0
  for x in zero['datasets']['stock']:x['opening']=0
 else:
  change['datasets']['tickets'][1]['elapsed']=10;change['datasets']['capacity'][1]['people']=2
  for x in zero['datasets']['tickets']:x['state']='취소'
 return [('base',c),('input_change',change),('zero_activity',zero)]
def convert(file,out,fmt):
 out.mkdir(parents=True,exist_ok=True);target=out/(file.stem+'.'+fmt)
 if target.exists():target.unlink()
 with tempfile.TemporaryDirectory(prefix='nubi-dual-lo-',dir='/private/tmp') as profile:
  result=subprocess.run([str(BIN/'soffice'),'-env:UserInstallation='+Path(profile).as_uri(),'--headless','--convert-to',fmt,'--outdir',str(out),str(file)],capture_output=True,timeout=90)
 if result.returncode or not target.exists():raise RuntimeError('conversion_failed')
 return target
def mutate(source,target,plan,case,variant):
 changes={}
 for sheet in plan['sheets']:
  for table in sheet.get('tables',[]):
   for ci,col in enumerate(table['columns']):
    if 'field' not in col:continue
    for ri,record in enumerate(variant['datasets'][table['dataset']]):
     value=record[col['field']];before=case['datasets'][table['dataset']][ri][col['field']]
     if value!=before:changes.setdefault(sheet['name'],{})[chr(64+table['col']+ci)+str(table['row']+1+ri)]=value
 if not changes:raise RuntimeError('no_meaningful_bound_mutation')
 with ZipFile(source) as z:
  wb=E.fromstring(z.read('xl/workbook.xml'));rels={x.get('Id'):x.get('Target') for x in E.fromstring(z.read('xl/_rels/workbook.xml.rels'))};parts={}
  for x in wb.find('s:sheets',NS):
   part=rels[x.get('{'+NS['r']+'}id')];parts[part.lstrip('/') if part.startswith('/') else 'xl/'+part]=changes.get(x.get('name'),{})
  with ZipFile(target,'w',ZIP_DEFLATED) as out:
   for name in z.namelist():
    raw=z.read(name)
    if re.fullmatch(r'xl/worksheets/sheet\d+\.xml',name):
     tree=E.fromstring(raw)
     for cell in tree.findall('.//s:c',NS):
      if cell.get('r') in parts.get(name,{}):
       value=parts[name][cell.get('r')]
       for x in list(cell):cell.remove(x)
       cell.attrib.pop('t',None)
       if isinstance(value,str):cell.set('t','inlineStr');E.SubElement(E.SubElement(cell,'{'+NS['s']+'}is'),'{'+NS['s']+'}t').text=value
       else:E.SubElement(cell,'{'+NS['s']+'}v').text=str(value)
      if cell.find('s:f',NS) is not None:
       for v in cell.findall('s:v',NS):cell.remove(v)
     raw=E.tostring(tree,encoding='utf-8',xml_declaration=True)
    out.writestr(name,raw)
def run(task):
 directory=ROOT/task['id'];source=ROOT/task['file'];plan=json.loads((directory/'composition.json').read_text());case=next(c for c in cases if c['id']==task['case']);row={'id':task['id'],'artifactHash':__import__('hashlib').sha256(source.read_bytes()).hexdigest(),'calculationCases':[],'status':'running'};out=QA/task['id'];out.mkdir(exist_ok=True)
 try:
  if task['domain']=='excel':
   bound={(t['dataset'],c['field']) for s in plan['sheets'] for t in s.get('tables',[]) for c in t['columns'] if 'field' in c};required={(name,field) for name,records in case['datasets'].items() for record in records for field in record};row['missingSourceFields']=[list(x) for x in sorted(required-bound)]
   for name,data in variants(case):
    work=out/name;work.mkdir(exist_ok=True);file=source
    if name!='base':file=work/'input.xlsx';mutate(source,file,plan,case,data)
    calculated=convert(file,work/'calculated','xlsx');values=load_workbook(calculated,data_only=True);formulas=load_workbook(calculated,data_only=False);checks=[]
    for metric,wanted in expected(data).items():
     location=next((x for x in plan['outputs'] if x['metric'] in [metric,case['metrics'][metric]]),None);actual=values[location['sheet']][location['cell']].value if location else None;formula=formulas[location['sheet']][location['cell']].data_type=='f' if location else False;tol=.00005 if metric=='sla' else .005 if metric in ['vat','reserve','total','gap','credit','potential'] and case['id']!='inventory' else 1e-7;ok=isinstance(actual,(int,float)) and not isinstance(actual,bool) and abs(actual-wanted)<=tol and formula;checks.append(dict(metric=metric,actual=actual,expected=wanted,calculatedCell=formula,passed=ok,tolerance=tol))
    raw=[]
    for s in plan['sheets']:
     for t in s.get('tables',[]):
      for ci,col in enumerate(t['columns']):
       if 'field' in col:
        for ri,record in enumerate(data['datasets'][t['dataset']]):raw.append(values[s['name']].cell(t['row']+ri+1,t['col']+ci).value==record[col['field']])
    errors=[f'{s.title}!{c.coordinate}' for s in values for cells in s for c in cells if c.data_type=='e'];row['calculationCases'].append(dict(variant=name,checks=checks,rawCellsChecked=len(raw),rawPreserved=all(raw),errors=errors,passed=all(x['passed'] for x in checks) and all(raw) and not errors and not row['missingSourceFields']))
    if name=='base':row['recalculated']=str(calculated);render_source=calculated
   row['calculationPassed']=all(x['passed'] for x in row['calculationCases']);source=render_source
  if '--no-render' not in sys.argv:
   render=out/'render';render.mkdir(exist_ok=True)
   for old in render.glob('page-*.png'):old.unlink()
   if task['domain']=='word':
    script=Path('/Users/ganadi/.codex/plugins/cache/openai-primary-runtime/documents/26.1009.11318/skills/documents/render_docx.py');subprocess.run([sys.executable,str(script),str(source),'--output_dir',str(render),'--emit_pdf','--dpi','85'],capture_output=True,check=True,timeout=90);pdf=render/(source.stem+'.pdf')
   else:
    pdf=convert(source,render,'pdf')
    for i,image in enumerate(convert_from_path(pdf,dpi=85,poppler_path=str(BIN)),1):image.save(render/f'page-{i}.png')
   pages=PdfReader(pdf).pages;row['pages']=len(pages);row['renderedText']=[p.extract_text() for p in pages];row['visualReviewRequired']=True
  row['status']='calculation_failed' if row.get('calculationPassed') is False else 'checked_pending_visual_review'
 except Exception as e:row['status']='failed';row['error']=type(e).__name__+': '+str(e)[:240]
 (out/'checks.json').write_text(json.dumps(row,ensure_ascii=False,indent=2)+'\n');print(json.dumps({k:v for k,v in row.items() if k in ['id','status','pages','calculationPassed','error']},ensure_ascii=False),flush=True)
if '--oracle-only' in sys.argv:
 print(json.dumps({c['id']:{name:expected(data) for name,data in variants(c)} for c in cases},ensure_ascii=False));sys.exit(0)
for p in sorted(ROOT.glob('*/attempt.json')):
 task=json.loads(p.read_text())
 if task.get('status')!='created_pending_independent_checks' or FILTER and FILTER not in task['id']:continue
 result=QA/task['id']/'checks.json'
 if result.exists() and '--force' not in sys.argv:continue
 run(task)
