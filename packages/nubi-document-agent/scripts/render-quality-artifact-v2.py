"""Host-side local rendering/readback for newly generated evaluation drafts.
Reuses bundled LibreOffice, Poppler, pdfplumber and openpyxl. No company reader.
"""
import sys,json,subprocess,tempfile,os,hashlib,re,unicodedata
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree as ET
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'eval'))
from pdf_text_evidence import text_evidence
import pdfplumber
from pdf2image import convert_from_path
from openpyxl import load_workbook
ROOT=Path(sys.argv[1]).resolve();cfg=json.loads((ROOT/'input.json').read_text());domain=cfg['domain'];plan=cfg['plan'];source=ROOT/('candidate.'+{'word':'docx','excel':'xlsx','powerpoint':'pptx'}[domain])
RUNTIME=Path('/Users/ganadi/.cache/codex-runtimes/codex-primary-runtime/dependencies');BIN=RUNTIME/'bin/override';font=ROOT/'fonts';font.mkdir(exist_ok=True);config=font/'fonts.conf';config.write_text(f'<fontconfig><dir>/System/Library/Fonts</dir><dir>/System/Library/Fonts/Supplemental</dir><cachedir>{font}/cache</cachedir></fontconfig>');os.environ.update(FONTCONFIG_FILE=str(config),SAL_FONTPATH='/System/Library/Fonts',TMPDIR='/private/tmp',PATH=str(BIN)+os.pathsep+os.environ.get('PATH',''))
sha=lambda b:hashlib.sha256(b).hexdigest();normalize=lambda t:re.sub(r'\s+','',unicodedata.normalize('NFC',str(t)))
def convert(file,out,fmt):
 out.mkdir(exist_ok=True,parents=True)
 with tempfile.TemporaryDirectory(prefix='nubi-quality-render-',dir='/private/tmp') as profile:
  run=subprocess.run([str(BIN/'soffice'),'-env:UserInstallation='+Path(profile).as_uri(),'--headless','--convert-to',fmt,'--outdir',str(out),str(file)],capture_output=True,timeout=90)
 result=out/(file.stem+'.'+('xlsx' if fmt.startswith('xlsx') else 'pdf'))
 if run.returncode or not result.exists():raise RuntimeError('render_conversion_failed')
 return result

def business_expected(case,data):
 if case=='renewal':
  accepted=[x for x in data['contracts'] if x['state']=='승인'];ids={x['id'] for x in accepted};base=round(sum(x['monthly']*x['months']*(1-x['discount']) for x in accepted),2);usage=sum(x['quantity']*x['price'] for x in data['usage'] if x['contract'] in ids and x['state']=='확정');vat=round((base+usage)*data['settings'][0]['vat']+1e-10,2);total=round(base+usage+vat,2)
  return dict(base=base,usage=usage,vat=vat,total=total,gap=round(data['settings'][0]['budget']-total,2),pending=round(sum(x['monthly']*x['months']*(1-x['discount']) for x in data['contracts'] if x['state']=='보류'),2),uncertain=sum(x['quantity']*x['price'] for x in data['usage'] if x['contract'] in ids and x['state']=='미확정'),approved=len(accepted))
 accepted=[x for x in data['shipments'] if x['state']=='확정'];refund=sum(x['returned']*x['refundUnit'] for x in accepted);shipping=sum(x['reship']*x['shippingUnit'] for x in accepted);work=sum(x['minutes'] for x in accepted);capacity=sum(x['people']*x['minutes'] for x in data['staff'])
 return dict(delivered=sum(x['delivered'] for x in accepted),returned=sum(x['returned'] for x in accepted),refund=refund,shipping=shipping,total=refund+shipping,pending=sum(x['returned']*x['refundUnit']+x['reship']*x['shippingUnit'] for x in data['shipments'] if x['state']=='보류'),work=work,capacity=capacity,timeGap=capacity-work,budgetGap=data['settings'][0]['budget']-refund-shipping)
def mutation_file(path,data):
 ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'};changes={}
 for sheet in plan['sheets']:
  for table in sheet.get('tables',[]):
   for ci,col in enumerate(table['columns']):
    if 'field' in col:
     for ri,record in enumerate(data[table['dataset']]):changes.setdefault(sheet['name'],{})[chr(64+table['col']+ci)+str(table['row']+ri+1)]=record[col['field']]
 with ZipFile(source) as z:
  wb=ET.fromstring(z.read('xl/workbook.xml'));rels={r.get('Id'):r.get('Target') for r in ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))};parts={}
  for sheet in wb.findall('s:sheets/s:sheet',ns):
   target=rels[sheet.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')];part=target.lstrip('/') if target.startswith('/') else 'xl/'+target;parts[part]=changes.get(sheet.get('name'),{})
  with ZipFile(path,'w',ZIP_DEFLATED) as out:
   for name in z.namelist():
    raw=z.read(name)
    if name in parts:
     tree=ET.fromstring(raw)
     for cell in tree.findall('.//s:c',ns):
      if cell.get('r') in parts[name]:
       v=parts[name][cell.get('r')]
       for child in list(cell):cell.remove(child)
       cell.attrib.pop('t',None)
       if isinstance(v,str):cell.set('t','inlineStr');ET.SubElement(ET.SubElement(cell,'{'+ns['s']+'}is'),'{'+ns['s']+'}t').text=v
       else:ET.SubElement(cell,'{'+ns['s']+'}v').text=str(int(v) if isinstance(v,bool) else v)
      if cell.find('s:f',ns) is not None:
       for child in cell.findall('s:v',ns):cell.remove(child)
     raw=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
    out.writestr(name,raw)

report={'artifactHash':sha(source.read_bytes()),'domain':domain,'pages':[],'layoutFindings':[],'calculation':None,'rawPreserved':True}
if domain=='excel':
 calc=convert(source,ROOT/'calculated','xlsx:Calc MS Excel 2007 XML');values=load_workbook(calc,data_only=True);formulas=load_workbook(calc,data_only=False);outputs=[];raw=[];bound=set()
 for s in plan['sheets']:
  for table in s.get('tables',[]):
   for ci,col in enumerate(table['columns']):
    if 'field' in col:
     bound.add((table['dataset'],col['field']))
     for ri,record in enumerate(cfg['datasets'][table['dataset']]):raw.append(values[s['name']].cell(table['row']+ri+1,table['col']+ci).value==record[col['field']])
 for o in plan['outputs']:
  outputs.append({'metric':o['metric'],'sheet':o['sheet'],'cell':o['cell'],'value':values[o['sheet']][o['cell']].value,'formula':formulas[o['sheet']][o['cell']].data_type=='f'})
 missing=[f'{ds}.{field}' for ds,records in cfg['datasets'].items() for record in records[:1] for field in record if (ds,field) not in bound]
 errors=[s.title+'!'+c.coordinate for s in values for row in s for c in row if c.data_type=='e'];report['rawPreserved']=bool(raw) and all(raw) and not missing;report['calculation']={'outputs':outputs,'errors':errors,'missingSourceFields':missing,'rawCellsChecked':len(raw)}
 report['calculation']['variants']=[]
 for variant in [{'id':'base','datasets':cfg['datasets']},*cfg.get('variants',[])]:
  data=variant['datasets'];work=ROOT/variant['id'];work.mkdir(exist_ok=True)
  if variant['id']=='base':vs=values;fs=formulas
  else:
   mutated=work/'input.xlsx';mutation_file(mutated,data);calculated=convert(mutated,work/'calculated','xlsx:Calc MS Excel 2007 XML');vs=load_workbook(calculated,data_only=True);fs=load_workbook(calculated,data_only=False)
  expected=business_expected(cfg['caseId'],data);metrics=[]
  for metric,wanted in expected.items():
   loc=next((o for o in plan['outputs'] if o['metric'] in [metric,cfg['metricNames'][metric]]),None);actual=vs[loc['sheet']][loc['cell']].value if loc else None;isformula=bool(loc and fs[loc['sheet']][loc['cell']].data_type=='f');ok=isinstance(actual,(int,float)) and not isinstance(actual,bool) and abs(actual-wanted)<.005 and isformula;metrics.append(dict(metric=metric,expected=wanted,actual=actual,formula=isformula,passed=ok))
  errors=[sh.title+'!'+ce.coordinate for sh in vs for rr in sh for ce in rr if ce.data_type=='e'];report['calculation']['variants'].append(dict(id=variant['id'],checks=metrics,errors=errors,passed=all(x['passed'] for x in metrics) and not errors))
 source=calc
pdf=convert(source,ROOT/'render','pdf');images=convert_from_path(pdf,dpi=100,poppler_path=str(BIN));fulltext=[]
with pdfplumber.open(pdf) as doc:
 for index,(page,image) in enumerate(zip(doc.pages,images)):
  path=ROOT/'render'/f'page-{index+1}.png';image.save(path);text=page.extract_text(x_tolerance=2) or '';fulltext.append(text);report['pages'].append({'page':index+1,'image':str(path),'imageHash':sha(path.read_bytes()),'text':text,'width':page.width,'height':page.height})
  chars=page.chars
  if any(c['x0']<-.5 or c['x1']>page.width+.5 or c['top']<-.5 or c['bottom']>page.height+.5 for c in chars):report['layoutFindings'].append({'code':'page_bounds','kind':'layout','severity':'major','path':'/'+{'word':'blocks','excel':'sheets','powerpoint':'slides'}[domain]+('/'+str(index) if domain=='powerpoint' else ''),'message':f'Rendered page {index+1} has glyphs beyond its page bounds.'})
  if domain=='powerpoint' and index<len(plan['slides']):
   for ei,e in enumerate(plan['slides'][index]['elements']):
    if e['kind'] not in ['text','shape'] or not e.get('text'):continue
    state=text_evidence(chars,e)
    if state!='present':report['layoutFindings'].append({'code':'render_text_'+('overflow' if state=='overflow' else 'missing'),'kind':'layout','severity':'major','path':f'/slides/{index}/elements/{ei}','message':'Actual spatial PDF glyph evidence shows '+state+' text. Reflow without dropping required facts.'})
if domain=='word':
 for bi,b in enumerate(plan['blocks']):
  if b['kind']=='paragraph' and b.get('style') in ['heading1','heading2']:
   for page in report['pages']:
    text=normalize(page['text']);needle=normalize(b['text']);at=text.find(needle)
    if at>=0 and len(text[at+len(needle):])<12:report['layoutFindings'].append({'code':'orphan_heading','kind':'layout','severity':'major','path':'/blocks','message':f'Heading at block {bi} has no substantive body below it on rendered page {page["page"]}.'})
report['renderHash']=sha(pdf.read_bytes());report['renderedText']='\n'.join(fulltext);(ROOT/'readback.json').write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str)+'\n');print(json.dumps({'domain':domain,'pages':len(report['pages']),'layoutFindings':len(report['layoutFindings'])}))
