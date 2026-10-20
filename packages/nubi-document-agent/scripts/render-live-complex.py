"""Reuse the established renderer; independently check new sensitivity outputs."""
import json,subprocess,sys
from pathlib import Path
from openpyxl import load_workbook

root=Path(sys.argv[1]).resolve()
subprocess.run([sys.executable,str(Path(__file__).with_name('render-quality-artifact-v2.py')),str(root)],check=True)
cfg=json.loads((root/'input.json').read_text())
if cfg['domain']=='excel' and any(o['metric']=='discountScenario' for o in cfg['plan']['outputs']):
    rb=json.loads((root/'readback.json').read_text())
    for state in [{'id':'base','datasets':cfg['datasets']},*cfg.get('variants',[])]:
        d=state['datasets'];active=[x for x in d['contracts'] if x['state']=='승인'];ids={x['id'] for x in active}
        usage=sum(x['quantity']*x['price'] for x in d['usage'] if x['contract'] in ids and x['state']=='확정')
        base=sum(x['monthly']*x['months']*(1-x['discount']) for x in active)
        reduced=sum(x['monthly']*x['months']*max(0,1-x['discount']-.05) for x in active)
        r=lambda n:round(n+1e-9,2)
        wants={'discountScenario':r(reduced+usage+r((reduced+usage)*d['settings'][0]['vat'])),'reducedBudgetGap':r(d['settings'][0]['budget']*.85-r(base+usage+r((base+usage)*d['settings'][0]['vat'])))}
        file=root/'calculated/candidate.xlsx' if state['id']=='base' else root/state['id']/'calculated/input.xlsx'
        values=load_workbook(file,data_only=True);formulas=load_workbook(file,data_only=False)
        checks=[]
        for metric,wanted in wants.items():
            loc=next((x for x in cfg['plan']['outputs'] if x['metric']==metric),None)
            actual=values[loc['sheet']][loc['cell']].value if loc else None
            formula=bool(loc and formulas[loc['sheet']][loc['cell']].data_type=='f')
            passed=isinstance(actual,(int,float)) and not isinstance(actual,bool) and abs(actual-wanted)<.005 and formula
            checks.append(dict(metric=metric,expected=wanted,actual=actual,formula=formula,passed=passed))
        receipt=next(x for x in rb['calculation']['variants'] if x['id']==state['id'])
        receipt['checks'].extend(checks);receipt['passed']=receipt['passed'] and all(x['passed'] for x in checks)
    (root/'readback.json').write_text(json.dumps(rb,ensure_ascii=False,indent=2)+'\n')
