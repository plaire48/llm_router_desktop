import sys,json,os,copy,tempfile,subprocess,re,io
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree as E
from openpyxl import load_workbook
from pypdf import PdfReader
from pdf2image import convert_from_path
root=Path(sys.argv[1]).resolve();runtime=Path('/Users/ganadi/.cache/codex-runtimes/codex-primary-runtime/dependencies');bin_dir=runtime/'bin/override';qa=root/'qa';qa.mkdir(exist_ok=True);fonts=qa/'fonts';fonts.mkdir(exist_ok=True);config=fonts/'fonts.conf';config.write_text(f'<?xml version="1.0"?><fontconfig><dir>/System/Library/Fonts</dir><dir>/System/Library/Fonts/Supplemental</dir><cachedir>{fonts}/cache</cachedir></fontconfig>');os.environ.update(FONTCONFIG_FILE=str(config),SAL_FONTPATH='/System/Library/Fonts',TMPDIR='/private/tmp',PATH=str(bin_dir)+os.pathsep+os.environ.get('PATH',''))
report=json.loads((root/'model-results.json').read_text());result={'tasks':[],'nativeOfficeVerified':False}
ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
def convert(source,out,fmt):
 out.mkdir(parents=True,exist_ok=True)
 with tempfile.TemporaryDirectory(prefix='nubi-open-qa-',dir='/private/tmp') as profile:
  run=subprocess.run([str(bin_dir/'soffice'),'-env:UserInstallation='+Path(profile).as_uri(),'--headless','--convert-to',fmt,'--outdir',str(out),str(source)],capture_output=True,timeout=60)
 target=out/(source.stem+'.'+fmt)
 if run.returncode or not target.exists():raise RuntimeError('conversion_failed')
 return target
def expected(c):
 data=c['datasets'];round2=lambda v:round(v+1e-10,2)
 if c['id']=='procurement':
  rows=[r for r in data['orders'] if r['state']=='유효'];a=sum((r['delivered']-r['rejected'])*r['unitPrice'] for r in rows);return {'accepted':a,'vat':round2(a*.1),'retained':round2(a*.05),'payable':round2(a*1.05),'backlog':sum(r['ordered']-r['delivered'] for r in rows)}
 if c['id']=='training':
  rows=data['registrations'];confirmed=sum(r['state']=='확정' for r in rows);revenue=round2(sum(0 if r['state']=='대기' else r['fee']*(1-r['discount'])*(.1 if r['state']=='환불' else 1) for r in rows));cost=confirmed*2+sum(r['instructor']+r['room'] for r in data['courses'])+40;return {'revenue':revenue,'confirmed':confirmed,'refunds':sum(r['state']=='환불' for r in rows),'cost':cost,'balance':round2(revenue-cost),'seats':36-confirmed}
 duration=sum(r['duration'] for r in data['tasks']);recovery=sum(r['rollbackMinutes'] for r in data['systems']);base=sum(r['duration']*r['people']*r['unitRate'] for r in data['tasks'])+120;total=round2(base*1.15);return {'duration':duration,'recovery':recovery,'required':duration+recovery,'windowGap':300-duration-recovery,'baseCost':base,'totalCost':total,'budgetGap':round2(1200-total),'pending':sum(r['state']=='대기' for r in data['tasks'])}
def alter(source,target,plan,dataset,field,value,all_rows=False,row_index=0):
 with ZipFile(source) as z:
  wb=E.fromstring(z.read('xl/workbook.xml'));rels={r.get('Id'):r.get('Target') for r in E.fromstring(z.read('xl/_rels/workbook.xml.rels'))};parts={s.get('name'):('xl/'+rels[s.get('{'+ns['r']+'}id')]).replace('xl//xl/','xl/') for s in wb.find('s:sheets',ns)};changes={};done=False
  for sheet in plan['sheets']:
   for table in sheet.get('tables',[]):
    if table['dataset']!=dataset:continue
    for col_index,col in enumerate(table['columns']):
     if col.get('field')==field:
      count=len(next(c for c in cases if c['id']==current_case)['datasets'][dataset]) if all_rows else 1
      for i in (range(count) if all_rows else [row_index]):
       addr=chr(64+table['col']+col_index)+str(table['row']+1+i);changes.setdefault(parts[sheet['name']],{})[addr]=value;done=True
  if not done:raise RuntimeError('mutation_input_not_exposed')
  with ZipFile(target,'w',ZIP_DEFLATED) as out:
   for name in z.namelist():
    raw=z.read(name)
    if re.match(r'xl/worksheets/sheet\d+\.xml$',name):
     tree=E.fromstring(raw)
     for cell in tree.findall('.//s:c',ns):
      if cell.get('r') in changes.get(name,{}):
       v=changes[name][cell.get('r')]
       for child in list(cell):cell.remove(child)
       cell.attrib.pop('t',None);E.SubElement(cell,'{'+ns['s']+'}v').text=str(v)
      if cell.find('s:f',ns) is not None:
       v=cell.find('s:v',ns)
       if v is not None:cell.remove(v)
     raw=E.tostring(tree,encoding='utf-8',xml_declaration=True)
    out.writestr(name,raw)
case_root=Path(__file__).resolve().parent.parent
cases=json.loads(subprocess.run([str(runtime/'node/bin/node'),'--input-type=module','-e',"import{cases}from'./eval/open-composition-cases.mjs';console.log(JSON.stringify(cases))"],cwd=case_root,capture_output=True,text=True,check=True).stdout)
for task in report['tasks']:
 inferred=task['id']+'/candidate.'+{'word':'docx','excel':'xlsx','powerpoint':'pptx'}[task['domain']]
 if not (root/task.get('file',inferred)).exists():continue
 directory=root/task['id'];plan=json.loads((directory/'composition.json').read_text());case=next(c for c in cases if c['id']==task['case']);current_case=case['id'];row={'id':task['id'],'domain':task['domain'],'calculationCases':[]};result['tasks'].append(row)
 try:
  source=root/task.get('file',inferred)
  if task['domain']=='excel':
   if case['id']=='procurement':dataset,field,newvalue='orders','unitPrice',30
   elif case['id']=='training':dataset,field,newvalue='registrations','fee',30
   else:dataset,field,newvalue='tasks','duration',40
   for variant in ['base','input_change','zero_amounts']:
    data=copy.deepcopy(case);variant_dir=qa/task['id']/variant;variant_dir.mkdir(parents=True,exist_ok=True);file=source
    if variant!='base':
     value=newvalue if variant=='input_change' else 0;file=variant_dir/'input.xlsx';index=1 if case['id']=='procurement' else 0;alter(source,file,plan,dataset,field,value,variant=='zero_amounts',index)
     targets=data['datasets'][dataset] if variant=='zero_amounts' else data['datasets'][dataset][index:index+1]
     for r in targets:r[field]=value
    output=convert(file,variant_dir/'calculated','xlsx');book=load_workbook(output,data_only=True);formula_book=load_workbook(output,data_only=False);want=expected(data);checks=[]
    for metric,value in want.items():
     loc=next((o for o in plan['outputs'] if o['metric'] in [metric,case['metrics'][metric]]),None);actual=book[loc['sheet']][loc['cell']].value if loc else None;formula=formula_book[loc['sheet']][loc['cell']].data_type=='f' if loc else False;checks.append({'metric':metric,'actual':actual,'expected':value,'calculatedCell':formula,'passed':isinstance(actual,(int,float)) and not isinstance(actual,bool) and abs(actual-value)<1e-7 and formula})
    raw_checks=[]
    for sheet in plan['sheets']:
     for table in sheet.get('tables',[]):
      for ci,col in enumerate(table['columns']):
       if 'field' not in col:continue
       for ri,record in enumerate(data['datasets'][table['dataset']]):
        got=book[sheet['name']].cell(table['row']+ri+1,table['col']+ci).value;wanted=record[col['field']];raw_checks.append(got==wanted)
    errors=[f'{s.title}!{c.coordinate} {c.value}' for s in book for row_cells in s for c in row_cells if c.data_type=='e'];row['calculationCases'].append({'variant':variant,'checks':checks,'errors':errors,'rawCellsChecked':len(raw_checks),'rawPreserved':all(raw_checks),'meaningfulInputChange':variant=='base' or want!=expected(case),'passed':not errors and all(c['passed'] for c in checks) and all(raw_checks) and (variant=='base' or want!=expected(case))})
    if variant=='base':row['recalculated']=str(output);source=output
   row['calculationPassed']=all(c['passed'] for c in row['calculationCases'])
  if '--no-render' not in sys.argv:
   out=qa/task['id']/'render';out.mkdir(parents=True,exist_ok=True)
   for stale in out.glob('page-*.png'):
    if re.fullmatch(r'page-\d+\.png',stale.name):stale.unlink()
   if task['domain']=='word':
    skill=Path('/Users/ganadi/.codex/plugins/cache/openai-primary-runtime/documents/26.1009.11318/skills/documents');subprocess.run([sys.executable,str(skill/'render_docx.py'),str(source),'--output_dir',str(out),'--emit_pdf','--dpi','100'],capture_output=True,check=True,timeout=60);pdf=out/(source.stem+'.pdf')
   else:
    pdf=convert(source,out,'pdf')
    for i,im in enumerate(convert_from_path(pdf,dpi=100,poppler_path=str(bin_dir)),1):im.save(out/f'page-{i}.png')
   pages=PdfReader(pdf).pages;row['pages']=len(pages);row['renderedText']=[p.extract_text() for p in pages];row['visualReviewRequired']=True
  row['status']='calculation_failed' if row.get('calculationPassed') is False else 'rendered_pending_review'
 except Exception as error:row['status']='failed';row['error']=str(error)
 (root/'qa-results.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n');print(json.dumps({k:v for k,v in row.items() if k in ['id','status','error','pages','calculationPassed']},ensure_ascii=False),flush=True)
