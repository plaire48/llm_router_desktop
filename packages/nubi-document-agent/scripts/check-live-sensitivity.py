"""Recalculate actual GenOffice edits and compare independent business arithmetic."""
import json,sys,subprocess,tempfile,os,hashlib
from pathlib import Path
from decimal import Decimal,ROUND_HALF_UP
from openpyxl import load_workbook

root=Path(sys.argv[1]).resolve();contract=json.loads((root/'contract.json').read_text());checks=[]
binary=Path('/Users/ganadi/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice')
D=lambda n:Decimal(str(n))
money=lambda n:n.quantize(D('.01'),rounding=ROUND_HALF_UP)
for row in contract['rows']:
 file=Path(row['file']);assert hashlib.sha256(file.read_bytes()).hexdigest()==row['editedHash']
 out=file.parent/'calculated';out.mkdir(exist_ok=True)
 with tempfile.TemporaryDirectory(prefix='nubi-sensitivity-',dir='/private/tmp') as profile:
  run=subprocess.run([str(binary),'-env:UserInstallation='+Path(profile).as_uri(),'--headless','--convert-to','xlsx:Calc MS Excel 2007 XML','--outdir',str(out),str(file)],capture_output=True,timeout=90,env={**os.environ,'TMPDIR':'/private/tmp'})
  assert run.returncode==0,run.stderr.decode(errors='replace')
 values=load_workbook(out/file.name,data_only=True);formulas=load_workbook(out/file.name,data_only=False)
 data=row['datasets'];active=[c for c in data['contracts'] if c['state']=='승인'];ids={c['id'] for c in active}
 usage=sum((D(u['quantity'])*D(u['price']) for u in data['usage'] if u['contract'] in ids and u['state']=='확정'),D(0))
 base=sum((D(c['monthly'])*D(c['months'])*(1-D(c['discount'])) for c in active),D(0))
 reduced=sum((D(c['monthly'])*D(c['months'])*max(D(0),1-D(c['discount'])-D(row['discount'])) for c in active),D(0))
 vat=D(data['settings'][0]['vat']);total=money(base+usage+money((base+usage)*vat))
 wanted={'discountScenario':money(reduced+usage+money((reduced+usage)*vat)),'reducedBudgetGap':money(D(data['settings'][0]['budget'])*(1-D(row['reduction']))-total)}
 assert len(row['outputs'])==2
 errors=[f'{s.title}!{c.coordinate}:{c.value}' for s in values for cells in s for c in cells if c.data_type=='e']
 for loc in row['outputs']:
  actual=values[loc['sheet']][loc['cell']].value;formula=formulas[loc['sheet']][loc['cell']].data_type=='f';expected=wanted[loc['metric']]
  passed=isinstance(actual,(int,float)) and not isinstance(actual,bool) and abs(D(actual)-expected)<D('.005') and formula and not errors
  checks.append(dict(provider=row['provider'],state=row['id'],metric=loc['metric'],expected=float(expected),actual=actual,formulaPreserved=formula,errors=errors,passed=passed))
result=dict(total=len(checks),passed=sum(c['passed'] for c in checks),checks=checks,nonTargetCellsPreserved=all(r['nonTargetCellsPreserved'] for r in contract['rows']),actualOffice2016=False)
(root/'result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n');print(json.dumps({k:result[k] for k in ['total','passed','nonTargetCellsPreserved']}));assert result['total']>0 and result['total']==len(contract['rows'])*2 and result['passed']==result['total']
