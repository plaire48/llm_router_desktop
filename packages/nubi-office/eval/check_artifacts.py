"""Real XLSX calculation check; neither a model evaluation nor Windows validation."""
import copy
import json
import shutil
import subprocess
import tempfile
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parents[1]
corpus = json.loads((ROOT / 'eval/cases.json').read_text())
destination = ROOT / 'verification/workbooks'
destination.mkdir(parents=True, exist_ok=True)
soffice = shutil.which('soffice')
if not soffice:
    raise SystemExit('LibreOffice unavailable; no calculation result was produced.')
version = subprocess.run([soffice, '--version'], check=True, capture_output=True, text=True).stdout.strip()
results = []
with tempfile.TemporaryDirectory(prefix='nubi-workbook-check-') as temp:
    source, output = Path(temp) / 'input', Path(temp) / 'calculated'
    source.mkdir(); output.mkdir()
    for case in corpus['cases']:
        book = Workbook(); sheet = book.active; sheet.title = 'CaseSheet'
        for address, data in case['cells'].items():
            cell = sheet[address]; cell.value = data.get('formula', data.get('value'))
            if 'bold' in data:
                font = copy.copy(cell.font); font.bold = data['bold']; cell.font = font
            if 'numberFormat' in data: cell.number_format = data['numberFormat']
        book.save(source / (case['id'] + '-input.xlsx'))
        # This constructs known answers; it does NOT call an agent or score one.
        for check in case['checks']:
            if check['operator'] != 'equals': continue
            cell, value = sheet[check['address']], check['expected']
            if check['field'] == 'formula': cell.value = value
            elif check['field'] == 'value' and cell.data_type != 'f': cell.value = value
            elif check['field'] == 'bold':
                font = copy.copy(cell.font); font.bold = value; cell.font = font
            elif check['field'] == 'numberFormat': cell.number_format = value
            elif check['field'] == 'fillColor': cell.fill = PatternFill('solid', fgColor=value.lstrip('#'))
        book.save(source / (case['id'] + '-golden.xlsx'))
    subprocess.run([soffice, '-env:UserInstallation=' + (Path(temp) / 'profile').as_uri(), '--headless', '--convert-to', 'xlsx', '--outdir', str(output), *map(str, source.glob('*.xlsx'))], check=True, capture_output=True, text=True, timeout=60)
    for case in corpus['cases']:
        path = output / (case['id'] + '-golden.xlsx')
        values, formulas = load_workbook(path, data_only=True), load_workbook(path, data_only=False)
        checked = []
        for check in case['checks']:
            cell = formulas.active[check['address']]
            actual = {'value': values.active[check['address']].value, 'formula': cell.value, 'bold': cell.font.bold, 'numberFormat': cell.number_format}.get(check['field'])
            assert actual == check['expected'], (case['id'], check, actual)
            checked.append({'address': check['address'], 'field': check['field'], 'expected': check['expected'], 'actual': actual})
        # Confirm formulas survive a real file save/reopen/recalculation cycle.
        for address, data in case['cells'].items():
            if 'formula' in data: assert formulas.active[address].value == data['formula']
        results.append({'case': case['id'], 'fixtureChecks': checked})
    for path in output.glob('*.xlsx'): shutil.copy2(path, destination / path.name)
report = {'kind': 'golden-fixture-calculation-only', 'checkedAt': datetime.now(timezone.utc).isoformat(), 'engine': version, 'xlsxFiles': len(list(destination.glob('*.xlsx'))), 'cases': results, 'modelCalled': False, 'modelAccuracy': None, 'productionExcelPortExecuted': False, 'windowsExcel2016Verified': False}
(ROOT / 'verification/artifact-calculation.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
print(json.dumps({'kind': report['kind'], 'files': report['xlsxFiles'], 'modelAccuracy': None}))
